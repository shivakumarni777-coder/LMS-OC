from pptx import Presentation
from pptx.util import Emu
import zipfile, os

path = r"C:\Users\KAVI SHIVA\OneDrive\Desktop\LMS-OC-SEMINAR.pptx"
prs = Presentation(path)
SW, SH = prs.slide_width, prs.slide_height
print("slide size: %.3f x %.3f in" % (SW/914400, SH/914400))
print("slides:", len(prs.slides._sldIdLst))
print("file size: %.1f KB" % (os.path.getsize(path)/1024))
print()
print("  #  shapes  pics  notes  title")
print("  -  ------  ----  -----  -----")
overflow = []
for i, s in enumerate(prs.slides, 1):
    pics = sum(1 for sh in s.shapes if sh.shape_type == 13)
    has_notes = s.has_notes_slide and bool(s.notes_slide.notes_text_frame.text.strip())
    title = ""
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            t = sh.text_frame.text.strip().split("\n")[0]
            if len(t) > 3 and t not in ("›",):
                title = t[:38]
                break
    for sh in s.shapes:
        if sh.left is None: continue
        r, b = sh.left + (sh.width or 0), sh.top + (sh.height or 0)
        if sh.left < -1000 or sh.top < -1000 or r > SW + 1000 or b > SH + 1000:
            overflow.append((i, sh.shape_type, round(sh.left/914400,2), round(sh.top/914400,2),
                             round(r/914400,2), round(b/914400,2)))
    print("  %-2d %6d  %4d  %5s  %s" % (i, len(s.shapes), pics, "yes" if has_notes else "NO", title))

print()
z = zipfile.ZipFile(path)
media = [n for n in z.namelist() if n.startswith("ppt/media/")]
print("embedded images:", len(media))
notes = [n for n in z.namelist() if "notesSlide" in n and n.endswith(".xml")]
print("notes slides:  ", len(notes))
print()
if overflow:
    print("OVERFLOW (%d shapes outside the slide):" % len(overflow))
    for o in overflow[:12]:
        print("   slide %-2d  l=%.2f t=%.2f r=%.2f b=%.2f" % (o[0],o[2],o[3],o[4],o[5]))
else:
    print("no shapes outside the slide bounds")
